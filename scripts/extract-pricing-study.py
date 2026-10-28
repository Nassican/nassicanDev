"""Read official DANE workbooks without editing them; retain source cells and hashes."""
import hashlib
import json
import os
from pathlib import Path
import openpyxl

directory = Path(os.environ["TEMP"]) / "nassican-pricing-study"
sources = [
    ("departamental2025.xlsx", "https://www.dane.gov.co/files/operaciones/PM/anex-PMDepartamental-2025.xlsx", "2026-08-12", 12, 35, 6),
    ("nacional2025.xlsx", "https://www.dane.gov.co/files/operaciones/PM/anex-PM-TotalNacional-2025.xlsx", "2026-06-12", 19, 41, 15),
]
result = {"year": 2025, "nationalIncome": 1281590, "metric": "Ingreso per cápita mensual de la unidad de gasto, pesos corrientes", "sources": [], "departments": [], "cities": []}
for filename, url, published, first, last, column in sources:
    path = directory / filename
    sheet = openpyxl.load_workbook(path, data_only=True, read_only=True)["Ingpc Act.Met."]
    assert sheet.cell(11 if column == 6 else 12, column).value == 2025
    if column == 6:
        assert sheet.cell(36, column).value == result["nationalIncome"]
    else:
        assert sheet.cell(13, column).value == result["nationalIncome"]
    result["sources"].append({"url": url, "published": published, "sheet": sheet.title, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
    for index in range(first, last + 1):
        name, income = sheet.cell(index, 1).value, sheet.cell(index, column).value
        assert isinstance(name, str) and isinstance(income, (int, float)) and income > 0
        result["departments" if column == 6 else "cities"].append({"name": name, "income": income, "cell": sheet.cell(index, column).coordinate})
output = Path(__file__).parent / "data" / "pricing-dane-2025.json"
output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(f"Extracted {len(result['departments'])} departments and {len(result['cities'])} cities; both national totals agree.")
